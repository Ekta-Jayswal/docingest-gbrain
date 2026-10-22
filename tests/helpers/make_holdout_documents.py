# What this file does: writes the held-out answer-key documents that are not plain markdown: a Word
# BRD and a PDF SOW. They are for measuring only, never for tuning. Change one only together with
# its answer key.

"""Write the held-out answer-key documents that are not plain markdown.

Usage: python tests/helpers/make_holdout_documents.py   (needs python-docx and reportlab)

  answer_keys/holdout/permit_portal_brd.docx      BRD, Word, outline from public BRD guides
  answer_keys/holdout/accessibility_audit_sow.pdf  SOW, PDF, outline of a government SOW

The held-out documents are for measuring, never for tuning: see answer_keys/holdout/README.md.
Change a document only together with its answer key.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer, Table

# Folder for the held-out documents: answer_keys/holdout.
OUT = Path(__file__).resolve().parents[2] / "answer_keys" / "holdout"


# ---- BRD (Word) ------------------------------------------------------------------------


# Add a Word table with a header row.
def word_table(doc, header: list[str], rows: list[list[str]]) -> None:
    """A Word table with a header row."""
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header, strict=True):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row, strict=True):
            cell.text = text


# Add a Word list, bulleted or numbered.
def word_list(doc, items: list[str], numbered: bool = False) -> None:
    """Bulleted or numbered list paragraphs."""
    for item in items:
        doc.add_paragraph(item, style="List Number" if numbered else "List Bullet")


# Build the permit portal BRD as a Word file, with its sections, tables and lists.
def permit_portal_brd() -> None:
    """Business requirements for an online building permit portal."""
    doc = Document()
    doc.add_heading("Business Requirements Document", level=0)
    doc.add_paragraph("Online Building Permit Applications", style="Subtitle")
    doc.add_paragraph("City of Maple Falls · Department of Planning and Building")

    doc.add_heading("Document Information", level=1)
    word_table(
        doc,
        ["Item", "Detail"],
        [
            ["Document ID", "BRD-PB-2026-04"],
            ["Version", "1.2"],
            ["Prepared by", "Priya Natarajan, Business Analyst"],
            ["Business owner", "Thomas Reyes, Chief Building Official"],
            ["Date", "2026-09-14"],
        ],
    )
    doc.add_heading("Revision History", level=2)
    word_table(
        doc,
        ["Version", "Date", "Author", "Changes"],
        [
            ["1.0", "2026-07-30", "Priya Natarajan", "First draft for review"],
            ["1.1", "2026-08-21", "Priya Natarajan", "Added inspection booking requirements"],
            ["1.2", "2026-09-14", "Priya Natarajan", "Updated after steering group review"],
        ],
    )

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "The City of Maple Falls issues about 4,600 building permits a year. Applicants "
        "submit paper forms at the permit counter or email scanned PDFs, and counter staff "
        "re-key every application into PermitTrack. A residential permit takes 41 days on "
        "average from application to decision, and applicants phone the office to book each "
        "inspection. This project will let applicants apply, pay and track their permits "
        "online, and let plans examiners and inspectors work from the same record. The "
        "department expects to halve residential turnaround and free about two counter "
        "positions for applicant help."
    )

    doc.add_heading("Project Objectives", level=1)
    word_list(
        doc,
        [
            "Cut the average time from application to decision for residential permits "
            "from 41 days to 20 days by December 2027.",
            "Receive at least 80% of permit applications online within twelve months of launch.",
            "Remove manual re-keying of application data by counter staff.",
            "Let applicants book inspections without phoning the office.",
        ],
        numbered=True,
    )

    doc.add_heading("Project Scope", level=1)
    doc.add_paragraph("The project includes:")
    word_list(
        doc,
        [
            "Online application forms for residential and small commercial building permits",
            "Online payment of permit fees",
            "Electronic plan review with comments returned to the applicant",
            "Inspection booking, and inspection results recorded in the field",
        ],
    )
    doc.add_paragraph("The project does not include:")
    word_list(
        doc,
        [
            "Planning and zoning applications, which stay on the current process",
            "Changes to the permit fee schedule",
            "Migration of permits closed before 2020",
        ],
    )
    doc.add_paragraph(
        "Phase 1 (residential permits) goes live in May 2027. Phase 2 (small commercial "
        "permits) follows in October 2027."
    )

    doc.add_heading("Key Stakeholders", level=1)
    word_table(
        doc,
        ["Name", "Role", "Involvement", "How affected"],
        [
            [
                "Thomas Reyes",
                "Chief Building Official",
                "Business owner; approves scope and budget",
                "Accountable for permit turnaround",
            ],
            [
                "Priya Natarajan",
                "Business Analyst, IT Services",
                "Writes and maintains the requirements",
                "None",
            ],
            [
                "Elena Vasquez",
                "Permit Counter Supervisor",
                "Reviews forms; daily user",
                "Counter work moves from data entry to applicant help",
            ],
            [
                "Marcus Hill",
                "Senior Plans Examiner",
                "Daily user of plan review",
                "Reviews drawings on screen instead of on paper",
            ],
            [
                "Joan Whitfield",
                "Lead Building Inspector",
                "Uses inspection booking in the field",
                "Receives inspection lists on a tablet",
            ],
            [
                "Contractors and homeowners",
                "Applicants",
                "Consulted through two workshops",
                "Apply, pay and book inspections online",
            ],
        ],
    )

    doc.add_heading("Current Process", level=1)
    doc.add_paragraph("A residential permit is handled today as follows:")
    word_list(
        doc,
        [
            "The applicant collects a paper form from the counter or downloads a PDF.",
            "Counter staff check the form for completeness and take payment by card or cheque.",
            "Counter staff type the application into PermitTrack.",
            "A plans examiner reviews the paper drawings and writes comments on them.",
            "The applicant phones the office to book each inspection.",
        ],
        numbered=True,
    )

    doc.add_heading("Business Requirements", level=1)
    doc.add_paragraph("The Objective column refers to the numbered project objectives.")
    word_table(
        doc,
        ["ID", "Requirement", "Priority", "Objective"],
        [
            [
                "BR-01",
                "Applicants can submit a complete permit application online, including drawings.",
                "Must have",
                "1, 2",
            ],
            [
                "BR-02",
                "Applicants can pay permit fees online at the time of application.",
                "Must have",
                "2",
            ],
            [
                "BR-03",
                "Application data reaches the permit record without re-keying.",
                "Must have",
                "3",
            ],
            [
                "BR-04",
                "Plans examiners can review drawings and return comments electronically.",
                "Must have",
                "1",
            ],
            [
                "BR-05",
                "Applicants can see the status of their application at any time.",
                "Nice to have",
                "1",
            ],
            [
                "BR-06",
                "Applicants can book, move and cancel inspections online.",
                "Must have",
                "4",
            ],
            [
                "BR-07",
                "Inspectors record results in the field and the applicant hears the same day.",
                "Nice to have",
                "4",
            ],
        ],
    )

    doc.add_heading("Functional Requirements", level=1)
    doc.add_paragraph(
        "Functional requirements are grouped by user. Each one traces to a business "
        "requirement, shown in brackets."
    )
    doc.add_heading("Applicant", level=2)
    word_list(
        doc,
        [
            "FR-1.1 The system shall let an applicant create an account with an email "
            "address and password. [BR-01]",
            "FR-1.2 The system shall show only the form sections that apply to the selected "
            "permit type. [BR-01]",
            "FR-1.3 The system shall accept drawings as PDF files of up to 50 MB each. [BR-01]",
            "FR-1.4 The system shall take card payment through CityPay before the application "
            "is submitted. [BR-02]",
            "FR-1.5 The system shall show the current status and the name of the assigned "
            "examiner. [BR-05]",
        ],
    )
    doc.add_heading("Plans Examiner", level=2)
    word_list(
        doc,
        [
            "FR-2.1 The system shall place a new application in the review queue within five "
            "minutes of payment. [BR-03]",
            "FR-2.2 The system shall let examiners mark up drawings and return comments to the "
            "applicant. [BR-04]",
            "FR-2.3 The system shall write the decision and its conditions to PermitTrack. "
            "[BR-03]",
        ],
    )
    doc.add_heading("Inspector", level=2)
    word_list(
        doc,
        [
            "FR-3.1 The system shall let applicants pick an inspection slot from the "
            "inspectors' open calendar. [BR-06]",
            "FR-3.2 The system shall send each inspector the day's inspections in the "
            "InspectorGo tablet app. [BR-07]",
            "FR-3.3 The system shall email the result to the applicant when the inspector "
            "closes the inspection. [BR-07]",
        ],
    )
    doc.add_heading("Business Rules", level=2)
    word_table(
        doc,
        ["Rule", "Description"],
        [
            ["BRL-01", "A deck more than 60 cm above ground needs a structural drawing."],
            [
                "BRL-02",
                "An application is complete only when every required document is attached "
                "and the fee is paid.",
            ],
            ["BRL-03", "Inspections must be booked by 3 pm on the previous working day."],
        ],
    )

    doc.add_heading("Data Requirements", level=1)
    word_table(
        doc,
        ["Data", "Source", "Retention"],
        [
            ["Permit application", "Applicant", "10 years after final inspection"],
            ["Drawings", "Applicant", "Life of the building"],
            ["Payment record", "CityPay", "7 years"],
            ["Inspection result", "Inspector", "10 years after final inspection"],
        ],
    )

    doc.add_heading("Non-Functional Requirements", level=1)
    doc.add_heading("Security", level=2)
    word_list(
        doc,
        [
            "NFR-01 Card data shall be handled only by CityPay; the portal shall not store "
            "card numbers.",
            "NFR-02 Staff shall sign in with their city network account.",
        ],
    )
    doc.add_heading("Availability", level=2)
    word_list(
        doc,
        [
            "NFR-03 The portal shall be available 99.5% of each month outside the Sunday "
            "01:00 to 05:00 maintenance window.",
        ],
    )
    doc.add_heading("Usability", level=2)
    word_list(
        doc,
        [
            "NFR-04 Every public page shall meet WCAG 2.2 level AA.",
            "NFR-05 The application form shall be usable on a phone.",
        ],
    )
    doc.add_heading("Performance", level=2)
    word_list(
        doc,
        ["NFR-06 A 50 MB drawing shall upload in under two minutes on a 10 Mbit/s connection."],
    )

    doc.add_heading("Interface Requirements", level=1)
    word_table(
        doc,
        ["Connected system", "Direction", "Data exchanged"],
        [
            ["PermitTrack", "Both ways", "Permit records, decisions and conditions"],
            ["CityPay", "Outbound", "Fee amounts and payment confirmations"],
            ["InspectorGo", "Both ways", "Inspection lists and results"],
            ["Parcel Registry", "Inbound", "Parcel number, zoning and owner name"],
        ],
    )

    doc.add_heading("Project Constraints", level=1)
    word_list(
        doc,
        [
            "Residential permits must be live before the counter lease ends in June 2027.",
            "The capital budget is capped at $650,000.",
            "PermitTrack stays the system of record and cannot be replaced in this project.",
        ],
    )

    doc.add_heading("Assumptions and Dependencies", level=1)
    doc.add_heading("Assumptions", level=2)
    word_list(
        doc,
        [
            "Most applicants have an email address and can scan or photograph their drawings.",
            "The PermitTrack vendor provides an integration interface at no extra cost.",
        ],
    )
    doc.add_heading("Dependencies", level=2)
    word_list(
        doc,
        [
            "The CityPay upgrade to version 4, scheduled by Finance for March 2027.",
            "Mobile network coverage for inspector tablets in the rural north district.",
        ],
    )

    doc.add_heading("Cost-Benefit Analysis", level=1)
    word_table(
        doc,
        ["Item", "Type", "Amount (USD)"],
        [
            ["Portal software licence, first three years", "Cost", "180,000"],
            ["Implementation and integration services", "Cost", "390,000"],
            ["Staff training", "Cost", "35,000"],
            ["Counter staff time no longer spent re-keying", "Benefit", "95,000 per year"],
            ["Printing and courier of paper drawings", "Benefit", "22,000 per year"],
        ],
    )
    doc.add_paragraph(
        "The $605,000 one-off cost is recovered in a little over five years from the yearly "
        "benefits alone, before counting faster permits for applicants."
    )

    doc.add_heading("Glossary", level=1)
    word_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Plans examiner", "A staff member who checks drawings against the building code."],
            ["Permit record", "The PermitTrack entry that holds a permit's full history."],
            ["Final inspection", "The last inspection, after which the permit is closed."],
        ],
    )

    doc.add_heading("Approval", level=1)
    word_table(
        doc,
        ["Name", "Title", "Signature", "Date"],
        [
            ["Thomas Reyes", "Chief Building Official", "Thomas Reyes", "2026-09-18"],
            ["Sandra Cole", "Director of IT Services", "", ""],
        ],
    )

    doc.save(OUT / "permit_portal_brd.docx")


# ---- SOW (PDF) -------------------------------------------------------------------------


# Build the accessibility audit SOW as a PDF, with its headings, paragraphs, lists and tables.
def accessibility_audit_sow() -> None:
    """A state agency SOW for a website accessibility audit, written as a PDF."""
    ss = getSampleStyleSheet()
    body, h1, h2 = ss["BodyText"], ss["Heading2"], ss["Heading3"]
    story: list = []

    # Add a paragraph to the PDF.
    def para(text: str) -> None:
        story.append(Paragraph(text, body))

    # Add a heading to the PDF, at the given level.
    def heading(text: str, level: int = 1) -> None:
        story.append(Paragraph(text, h1 if level == 1 else h2))

    # Add a bulleted list to the PDF, or a lettered one when asked.
    def items(texts: list[str], lettered: bool = False) -> None:
        story.append(
            ListFlowable(
                [Paragraph(t, body) for t in texts],
                bulletType="a" if lettered else "bullet",
                start="a" if lettered else None,
                leftIndent=14,
            )
        )

    # Add a grid table to the PDF, with the given column widths.
    def table(rows: list[list[str]], widths: list[float]) -> None:
        cells = [[Paragraph(c, body) for c in r] for r in rows]
        t = Table(cells, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
        story.append(t)
        story.append(Spacer(1, 4 * mm))

    story.append(Paragraph("STATEMENT OF WORK", ss["Title"]))
    story.append(Paragraph("Website Accessibility Audit and Remediation", ss["Heading2"]))
    para(
        "<b>Contract No.:</b> DOR-2026-0412<br/>"
        "<b>Agency:</b> Department of Revenue, State of Westmore<br/>"
        "<b>Contractor:</b> Clearpath Digital Inc.<br/>"
        "<b>Date:</b> 2026-10-12"
    )

    heading("1. Purpose")
    para(
        "The Department of Revenue (the Department) requires a contractor to audit its public "
        "website and online tax forms for accessibility, fix the failures found, and train "
        "Department staff to keep new content accessible. The Contractor shall provide all "
        "labour, tools and management needed to perform the work in this Statement of Work."
    )

    heading("2. Background")
    para(
        "The Department's website has about 1,900 pages and 85 fillable PDF forms used by "
        "individual and business taxpayers. In 2026 the Department received complaints from "
        "taxpayers who use screen readers and could not complete the estimated tax payment "
        "form. State law requires all agency websites to conform to the state accessibility "
        "standard by 31 December 2027."
    )

    heading("3. Scope of Work")
    para("The Contractor shall perform the following tasks.")
    heading("3.1 Task 1 - Accessibility Audit", 2)
    para("The Contractor shall:")
    items(
        [
            "Test a sample of 300 pages, chosen with the Department, against WCAG 2.2 level AA "
            "using automated tools and manual testing with two screen readers.",
            "Test all 85 fillable PDF forms.",
            "Record each failure with its WCAG success criterion, location and severity.",
        ],
        lettered=True,
    )
    heading("3.2 Task 2 - Remediation", 2)
    para("The Contractor shall:")
    items(
        [
            "Fix all critical and serious failures in the website templates and shared "
            "components.",
            "Rebuild the 40 most used PDF forms as accessible web forms.",
            "Re-test every fixed page and form and record the result.",
        ],
        lettered=True,
    )
    heading("3.3 Task 3 - Training and Handover", 2)
    para("The Contractor shall:")
    items(
        [
            "Run four half-day training sessions for Department content editors.",
            "Write an accessibility checklist for new content.",
        ],
        lettered=True,
    )

    heading("4. Location of Work")
    para(
        "The Contractor shall perform the work at its own premises. Training sessions shall "
        "be held on site at the Department's offices, 200 Capitol Avenue, Riverton."
    )

    heading("5. Period of Performance")
    para("The period of performance is from 4 January 2027 through 30 September 2027.")

    heading("6. Deliverables Schedule")
    table(
        [
            ["Deliverable No.", "Deliverable", "Task", "Due Date"],
            ["1", "Audit plan and page sample", "Task 1", "2027-01-29"],
            ["2", "Audit report with prioritized failures", "Task 1", "2027-03-12"],
            ["3", "Remediated templates and components", "Task 2", "2027-06-18"],
            ["4", "Accessible web forms (40)", "Task 2", "2027-08-13"],
            ["5", "Training sessions and editor checklist", "Task 3", "2027-09-10"],
            ["6", "Final conformance report", "Task 2", "2027-09-24"],
        ],
        [28, 82, 22, 30],
    )

    heading("7. Applicable Standards")
    para("The Contractor shall perform the work in accordance with:")
    items(
        [
            "Web Content Accessibility Guidelines (WCAG) 2.2, level AA",
            "Section 508 of the Rehabilitation Act, as amended",
            "Westmore IT Accessibility Standard ITS-560, version 3",
        ]
    )

    heading("8. Acceptance Criteria")
    para(
        "The Department will review each deliverable within ten business days of receipt and "
        "either accept it in writing or return it with the reasons for rejection. The "
        "Contractor shall correct a rejected deliverable within ten business days."
    )
    table(
        [
            ["Deliverable", "Acceptance criterion"],
            [
                "Audit report with prioritized failures",
                "Every failure lists the WCAG success criterion, the page and the severity.",
            ],
            [
                "Remediated templates and components",
                "No critical or serious failures remain when re-tested with the Department's "
                "tool and two screen readers.",
            ],
            [
                "Accessible web forms (40)",
                "Each form can be completed and submitted using only a keyboard.",
            ],
            [
                "Final conformance report",
                "Signed off by the Department's Accessibility Coordinator.",
            ],
        ],
        [62, 100],
    )

    heading("9. Key Personnel")
    para(
        "The Contractor's key personnel are listed below. The Contractor shall not replace "
        "key personnel without the Department's written approval."
    )
    table(
        [
            ["Name", "Position", "Minimum qualification"],
            ["Devon Price", "Project Manager", "Five years managing web projects"],
            ["Aisha Karim", "Lead Accessibility Tester", "IAAP Web Accessibility Specialist"],
            ["Leo Brandt", "Front-end Developer", "Three years building accessible web forms"],
        ],
        [40, 52, 70],
    )

    heading("10. Points of Contact")
    table(
        [
            ["Role", "Name", "Email", "Phone"],
            ["Contracting Officer", "Margaret Dunn", "m.dunn@dor.westmore.example", "555-0142"],
            ["Department Project Manager", "Kevin Ambrose", "k.ambrose@dor.westmore.example", "555-0177"],
            ["Accessibility Coordinator", "Lucia Ferreira", "l.ferreira@dor.westmore.example", "555-0165"],
            ["Contractor Project Manager", "Devon Price", "devon.price@clearpath.example", "555-0190"],
        ],
        [44, 32, 58, 28],
    )

    heading("11. Special Requirements")
    items(
        [
            "All Contractor staff with access to Department systems shall pass a state "
            "background check before starting work.",
            "The Contractor shall not copy taxpayer data out of the Department's test "
            "environment.",
            "Work on the production website shall be done only between 19:00 and 06:00.",
        ]
    )

    heading("12. Contract Type and Payment Schedule")
    para(
        "This is a firm-fixed-price contract with a total price of $268,500. The Department "
        "will pay each amount below within 30 days of accepting the deliverable it is tied to. "
        "Training is included in the final payment."
    )
    table(
        [
            ["Payment", "Deliverable", "Amount"],
            ["Payment 1", "Audit plan and page sample", "$21,500"],
            ["Payment 2", "Audit report with prioritized failures", "$64,000"],
            ["Payment 3", "Remediated templates and components", "$88,000"],
            ["Payment 4", "Accessible web forms (40)", "$62,000"],
            ["Payment 5", "Final conformance report", "$33,000"],
        ],
        [30, 100, 32],
    )

    heading("13. Miscellaneous")
    para(
        "Travel costs are not reimbursed. Any change to this Statement of Work requires a "
        "written amendment signed by the Contracting Officer. All content, code and reports "
        "produced under this Statement of Work belong to the Department."
    )

    heading("14. Signatures")
    para(
        "<b>For the Department of Revenue</b><br/>Signature: /s/ Margaret Dunn<br/>"
        "Name: Margaret Dunn<br/>Title: Contracting Officer<br/>Date: 2026-10-14"
    )
    story.append(Spacer(1, 4 * mm))
    para(
        "<b>For Clearpath Digital Inc.</b><br/>Signature: /s/ Rebecca Lin<br/>"
        "Name: Rebecca Lin<br/>Title: Chief Executive Officer<br/>Date: 2026-10-13"
    )

    SimpleDocTemplate(
        str(OUT / "accessibility_audit_sow.pdf"),
        pagesize=A4,
        title="Statement of Work: Website Accessibility Audit and Remediation",
        leftMargin=20 * mm,
        rightMargin=20 * mm,
    ).build(story)


# Build both held-out documents.
def main() -> None:
    """Build and save both documents."""
    OUT.mkdir(parents=True, exist_ok=True)
    permit_portal_brd()
    accessibility_audit_sow()
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
