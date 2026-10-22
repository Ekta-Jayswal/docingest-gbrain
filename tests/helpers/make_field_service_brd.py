# What this file does: writes the field-service BRD answer-key document, a Word file with real
# headings. Change the document only together with its answer key.

"""Write answer_keys/field_service_brd.docx: a synthetic BRD with real Word headings.

Usage: python tests/helpers/make_field_service_brd.py   (needs python-docx)

The document is an answer-key document: change it only together with
answer_keys/field_service_brd.yaml.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

# Where the generated document is saved: answer_keys/.
OUT = Path(__file__).resolve().parents[2] / "answer_keys" / "field_service_brd.docx"


# Add a Word table with a grid style, filled with the header row and the data rows.
def table(doc, header: list[str], rows: list[list[str]]) -> None:
    """A Word table with a header row."""
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header, strict=True):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row, strict=True):
            cell.text = text


# Build the document heading by heading, then save it.
def main() -> None:
    """Build and save the document."""
    doc = Document()
    doc.add_heading("Business Requirements Document: Field Service Scheduling", level=0)
    doc.add_paragraph("Riverbend Water Utility · Customer Operations")

    doc.add_heading("1. Document Control", level=1)
    table(
        doc,
        ["Field", "Value"],
        [
            ["Document ID", "BRD-FSS-007"],
            ["Version", "0.9"],
            ["Author", "Marta Kowalski"],
            ["Sponsor", "Daniel Osei"],
            ["Status", "Draft for review"],
        ],
    )

    doc.add_heading("2. Purpose", level=1)
    doc.add_paragraph(
        "Riverbend Water Utility books about 3,200 field visits a month for meter repairs, "
        "leak checks and new connections. Dispatchers schedule them by hand from a shared "
        "spreadsheet, so crews often arrive without the right parts and customers wait at "
        "home for half a day. This document sets out what the business needs from a "
        "scheduling solution; it does not choose one."
    )

    doc.add_heading("3. Business Objectives", level=1)
    table(
        doc,
        ["ID", "Objective", "Measure"],
        [
            ["OBJ-1", "Shorter appointment windows for customers", "Window of 2 hours or less"],
            ["OBJ-2", "First-visit fixes", "85% of jobs fixed on the first visit"],
            ["OBJ-3", "Less dispatcher effort", "Scheduling time per job cut by half"],
        ],
    )

    doc.add_heading("4. Stakeholders", level=1)
    table(
        doc,
        ["ID", "Name", "Title", "Department", "RACI"],
        [
            ["SH-01", "Daniel Osei", "Director of Customer Operations", "Operations", "A"],
            ["SH-02", "Marta Kowalski", "Business Analyst", "Transformation Office", "R"],
            ["SH-03", "Grace Liu", "Field Operations Manager", "Operations", "C"],
            ["SH-04", "Ahmed Rahimi", "IT Applications Lead", "Information Technology", "C"],
        ],
    )

    doc.add_heading("5. Business Requirements", level=1)
    table(
        doc,
        ["ID", "Requirement", "Objective", "Owner"],
        [
            [
                "BR-001",
                "Customers can choose a two-hour appointment window online.",
                "OBJ-1",
                "Grace Liu",
            ],
            [
                "BR-002",
                "Crews know the likely parts before they leave the depot.",
                "OBJ-2",
                "Grace Liu",
            ],
            [
                "BR-003",
                "Jobs are assigned to crews automatically by skill and location.",
                "OBJ-3",
                "Ahmed Rahimi",
            ],
            ["BR-004", "Customers are told when the crew is on its way.", "OBJ-1", "Daniel Osei"],
        ],
    )

    doc.add_heading("6. Functional Requirements", level=1)
    table(
        doc,
        ["ID", "Requirement", "Business Requirement", "System"],
        [
            [
                "FR-001",
                "Offer open two-hour windows on the customer website.",
                "BR-001",
                "CustomerHub",
            ],
            [
                "FR-002",
                "Reserve the window when the customer confirms it.",
                "BR-001",
                "WorkForce Planner",
            ],
            ["FR-003", "Suggest parts from the asset's repair history.", "BR-002", "AssetTrack"],
            [
                "FR-004",
                "Assign each job to the nearest crew with the required skill.",
                "BR-003",
                "WorkForce Planner",
            ],
            [
                "FR-005",
                "Rebalance the day's jobs when a crew calls in sick.",
                "BR-003",
                "WorkForce Planner",
            ],
            [
                "FR-006",
                "Send a text message when the crew leaves for the job.",
                "BR-004",
                "CustomerHub",
            ],
        ],
    )

    doc.add_heading("7. Systems", level=1)
    table(
        doc,
        ["ID", "System", "Owner"],
        [
            ["SYS-01", "WorkForce Planner", "Ahmed Rahimi"],
            ["SYS-02", "AssetTrack", "Ahmed Rahimi"],
            ["SYS-03", "CustomerHub", "Daniel Osei"],
        ],
    )

    doc.add_heading("8. Assumptions and Risks", level=1)
    doc.add_heading("8.1 Assumptions", level=2)
    table(
        doc,
        ["ID", "Assumption", "Related Requirements"],
        [
            [
                "AS-01",
                "AssetTrack holds at least two years of repair history per asset.",
                "FR-003",
            ],
            [
                "AS-02",
                "Crew locations are available from the vans' GPS units every minute.",
                "FR-004, FR-005",
            ],
        ],
    )
    doc.add_heading("8.2 Risks", level=2)
    table(
        doc,
        ["ID", "Risk", "Likelihood", "Impact", "Affected Requirements"],
        [
            ["RK-01", "Crews keep using the spreadsheet in parallel.", "Medium", "High", "BR-003"],
            [
                "RK-02",
                "Repair history is too sparse to suggest parts.",
                "High",
                "Medium",
                "BR-002, FR-003",
            ],
        ],
    )

    doc.add_heading("9. Glossary", level=1)
    table(
        doc,
        ["Term", "Definition"],
        [
            ["Appointment window", "The period in which the crew promises to arrive."],
            ["Crew", "One or two field technicians working from one van."],
            ["First-visit fix", "A job completed without a return visit."],
        ],
    )

    doc.add_heading("10. Approval", level=1)
    table(
        doc,
        ["Approver", "Role", "Decision", "Date"],
        [
            ["Daniel Osei", "Sponsor", "Approved", "2026-10-02"],
            ["Ahmed Rahimi", "IT Applications Lead", "Pending", ""],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
