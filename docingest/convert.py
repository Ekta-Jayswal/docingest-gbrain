"""Convert input files to markdown.

Each converter returns the markdown plus an independent plain-text
extraction of the source (the "reference text"). Verification compares the
two token by token to prove nothing was lost in conversion.
"""

from __future__ import annotations

import csv
import email
import io
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from email import policy
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path

from bs4 import BeautifulSoup
from markdownify import MarkdownConverter

from .mdutil import md_table, slugify
from .model import Converted

ROWS_PER_TABLE = 100

SUPPORTED = {
    ".md",
    ".markdown",
    ".txt",
    ".text",
    ".log",
    ".docx",
    ".doc",
    ".pdf",
    ".html",
    ".htm",
    ".csv",
    ".tsv",
    ".xlsx",
    ".xlsm",
    ".pptx",
    ".eml",
    ".msg",
    ".vtt",
    ".srt",
}


class ConversionError(Exception):
    """A file could not be converted to markdown."""


def convert(path: Path, attach_dir: Path) -> tuple[Converted, str]:
    """Convert one file to markdown.

    Returns the ``Converted`` result and an independent plain-text extraction
    of the source, used by verification to prove no words were lost.
    """
    ext = path.suffix.lower()
    handler = {
        ".md": _markdown,
        ".markdown": _markdown,
        ".txt": _text,
        ".text": _text,
        ".log": _text,
        ".docx": _docx,
        ".doc": _doc,
        ".pdf": _pdf,
        ".html": _html,
        ".htm": _html,
        ".csv": _csv,
        ".tsv": _csv,
        ".xlsx": _xlsx,
        ".xlsm": _xlsx,
        ".pptx": _pptx,
        ".eml": _eml,
        ".msg": _msg,
        ".vtt": _subtitles,
        ".srt": _subtitles,
    }.get(ext)
    if handler is None:
        raise ConversionError(f"Unsupported file type: {ext}")
    conv, ref = handler(path, attach_dir)
    md = (
        re.sub(r"\n{3,}", "\n\n", conv.markdown.replace("\r\n", "\n").replace("\r", "\n")).strip()
        + "\n"
    )
    conv.markdown = md
    return conv, ref


def read_text_any(path: Path) -> str:
    """Read a text file, trying UTF-8 (with BOM), then Windows-1252, then Latin-1."""
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


# ---- text and markdown ---------------------------------------------------------


def _markdown(path: Path, _attach: Path) -> tuple[Converted, str]:
    text = read_text_any(path)
    return Converted(markdown=text, converter="markdown"), text


DIALOGUE_PATTERNS = [
    # WhatsApp / chat exports: "12/01/2026, 10:15 - Name: msg" or "[12/01/26, 10:15:22] Name: msg"
    (
        "chat",
        re.compile(
            r"^\[?\d{1,4}[/.-]\d{1,2}[/.-]\d{2,4},?\s+\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AaPp][Mm])?\]?\s*(?:-\s*)?([^:\n]{1,40}?):\s"
        ),
    ),
    # Transcripts with timestamps: "[00:01:02] Name: text" or "00:01 Name: text"
    (
        "transcript",
        re.compile(r"^\[?\(?\d{1,2}:\d{2}(?::\d{2})?\)?\]?\s*-?\s*([A-Z][^:\n]{0,40}?):\s"),
    ),
    # Speaker-only transcripts: "Name: text"
    ("transcript", re.compile(r"^\**([A-Z][A-Za-z.'\- ]{1,40}?)\**:\s+\S")),
]


def detect_dialogue(text: str) -> tuple[str | None, list[str]]:
    """Detect chat or transcript text by its line pattern.

    Returns the kind (``chat``, ``transcript`` or None) and the speakers,
    most frequent first.
    """
    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) < 4:
        return None, []
    for kind, pattern in DIALOGUE_PATTERNS:
        speakers = [m.group(1).strip() for line in lines if (m := pattern.match(line))]
        if len(speakers) >= 0.6 * len(lines):
            counts = Counter(s for s in speakers if 1 <= len(s.split()) <= 4)
            return kind, [s for s, _ in counts.most_common()]
    return None, []


def _text(path: Path, _attach: Path) -> tuple[Converted, str]:
    text = read_text_any(path)
    kind, speakers = detect_dialogue(text)
    conv = Converted(markdown=text, converter="text", kind_hint=kind)
    if kind:
        conv.meta["participants"] = speakers
    return conv, text


# ---- HTML-based (HTML, DOCX) -------------------------------------------------------


class _Converter(MarkdownConverter):
    """markdownify with tables rendered as clean pipe tables (first row = header)."""

    def convert_table(self, el, text, *args, **kwargs):
        """Render an HTML table as a pipe table, ignoring nested tables."""
        rows = []
        for tr in el.find_all("tr"):
            if tr.find_parent("table") is not el:
                continue
            rows.append(
                [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"], recursive=False)]
            )
        rows = [r for r in rows if any(c for c in r)]
        if not rows:
            return ""
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        return "\n\n" + md_table(rows[0], rows[1:]) + "\n\n"


def html_to_markdown(html: str) -> str:
    """Convert HTML to markdown, dropping scripts and styles."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    return _Converter(
        heading_style="ATX", bullets="-", escape_underscores=False, escape_asterisks=False
    ).convert_soup(soup)


def _html(path: Path, _attach: Path) -> tuple[Converted, str]:
    html = read_text_any(path)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    root = soup.body or soup
    md = html_to_markdown(str(root))
    reference = root.get_text(" ")
    title = soup.title.get_text(strip=True) if soup.title else ""
    if title and not re.search(r"^# ", md, re.M):
        md = f"# {title}\n\n{md}"
        reference = f"{title} {reference}"
    return Converted(markdown=md, converter="html"), reference


def _docx(path: Path, _attach: Path) -> tuple[Converted, str]:
    import docx
    import mammoth

    with open(path, "rb") as f:
        result = mammoth.convert_to_html(f)
    conv = Converted(
        markdown=html_to_markdown(result.value),
        converter="docx (mammoth)",
        warnings=[m.message for m in result.messages if m.type == "error"],
    )
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return conv, "\n".join(parts)


def _doc(path: Path, attach: Path) -> tuple[Converted, str]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ConversionError(
            "Legacy .doc needs LibreOffice (soffice) to convert; install it "
            "(the Docker image includes it) or save the file as .docx"
        )
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)],
                check=True,
                capture_output=True,
                timeout=300,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            raise ConversionError(f"LibreOffice could not convert the file: {e}") from e
        out = Path(tmp) / f"{path.stem}.docx"
        if not out.exists():
            raise ConversionError("LibreOffice did not produce a .docx")
        conv, ref = _docx(out, attach)
    conv.converter = "doc (libreoffice + mammoth)"
    return conv, ref


# ---- PDF -------------------------------------------------------------------------


def _pdf(path: Path, _attach: Path) -> tuple[Converted, str]:
    import pdfplumber

    pages, size_chars, reference = [], Counter(), []
    with pdfplumber.open(str(path)) as pdf:
        n_pages = len(pdf.pages)
        for page in pdf.pages:
            reference.append(page.extract_text() or "")
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]

            def outside(obj, boxes=boxes):
                if obj.get("object_type") != "char":
                    return True
                x, y = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
                return not any(b[0] <= x <= b[2] and b[1] <= y <= b[3] for b in boxes)

            words = page.filter(outside).extract_words(
                extra_attrs=["size", "fontname"], use_text_flow=True
            )
            lines: list[dict] = []
            for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
                if lines and abs(w["top"] - lines[-1]["top"]) <= 2.5:
                    line = lines[-1]
                    line["words"].append(w)
                    line["bottom"] = max(line["bottom"], w["bottom"])
                else:
                    lines.append({"top": w["top"], "bottom": w["bottom"], "words": [w]})
            items = []
            for line in lines:
                ws = sorted(line["words"], key=lambda w: w["x0"])
                text = " ".join(w["text"] for w in ws)
                size = round(max(w["size"] for w in ws), 1)
                bold = all("bold" in w["fontname"].lower() for w in ws)
                size_chars[size] += len(text)
                items.append(("line", line["top"], line["bottom"], text, size, bold))
            for t in tables:
                rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in t.extract()]
                rows = [r for r in rows if any(r)]
                if rows:
                    items.append(("table", t.bbox[1], t.bbox[3], rows, 0, False))
            pages.append(sorted(items, key=lambda i: i[1]))

    conv = Converted(markdown="", converter="pdf (pdfplumber)", meta={"page_count": n_pages})
    total_chars = sum(size_chars.values())
    if total_chars < 30 * max(n_pages, 1):
        conv.warnings.append(
            f"PDF has almost no text layer ({total_chars} chars over {n_pages} pages); "
            "it is probably scanned and needs OCR"
        )
    body_size = size_chars.most_common(1)[0][0] if size_chars else 0
    heading_sizes = sorted(
        {s for s in size_chars if body_size and s >= body_size * 1.15}, reverse=True
    )[:3]
    level_of = {s: i + 1 for i, s in enumerate(heading_sizes)}
    bold_level = min(len(heading_sizes) + 1, 4)

    out: list[str] = []
    for items in pages:
        prev_bottom, prev_heading = None, None
        for kind, top, bottom, content, size, bold in items:
            if kind == "table":
                out += ["", md_table(content[0], content[1:]), ""]
                prev_bottom, prev_heading = bottom, None
                continue
            text = content
            level = level_of.get(size) if len(text) < 150 else None
            if (
                level is None
                and bold
                and len(text) < 100
                and not text.endswith((".", ",", ";"))
                and size >= body_size * 0.95
            ):
                level = bold_level
            gap = (top - prev_bottom) if prev_bottom is not None else 0
            if level:
                if prev_heading == (level, size) and gap < size:  # wrapped heading line
                    last = max(i for i, line in enumerate(out) if line.strip())
                    out[last] += " " + text
                else:
                    out += ["", "#" * level + " " + text, ""]
                prev_heading = (level, size)
            else:
                if prev_bottom is not None and gap > max(size, 1) * 0.6:
                    out.append("")
                out.append(text)
                prev_heading = None
            prev_bottom = bottom
        out.append("")
    conv.markdown = "\n".join(out)
    return conv, "\n".join(reference)


# ---- spreadsheets and slides ----------------------------------------------------------


def _rows_to_md(title: str, rows: list[list[str]], level: int) -> str:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return f"{'#' * level} {title}\n\n(empty)\n"
    header, body = rows[0], rows[1:]
    if len(body) <= ROWS_PER_TABLE:
        return f"{'#' * level} {title}\n\n{md_table(header, body)}\n"
    parts = [f"{'#' * level} {title}\n"]
    for i in range(0, len(body), ROWS_PER_TABLE):
        chunk = body[i : i + ROWS_PER_TABLE]
        parts.append(
            f"{'#' * (level + 1)} Rows {i + 1}–{i + len(chunk)}\n\n{md_table(header, chunk)}\n"
        )
    return "\n".join(parts)


def _csv(path: Path, _attach: Path) -> tuple[Converted, str]:
    text = read_text_any(path)
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    md = f"# {path.stem}\n\n" + _rows_to_md("Data", rows, 2)
    return Converted(
        markdown=md,
        converter="csv",
        kind_hint="spreadsheet",
        meta={"row_count": max(len(rows) - 1, 0)},
    ), "\n".join(" ".join(r) for r in rows)


def _cell(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def _xlsx(path: Path, _attach: Path) -> tuple[Converted, str]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts, ref = [f"# {path.stem}\n"], []
    try:  # read-only workbooks keep the file open until closed
        for ws in wb.worksheets:
            rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
            ref.extend(" ".join(r) for r in rows)
            parts.append(_rows_to_md(f"Sheet: {ws.title}", rows, 2))
        sheet_count = len(wb.worksheets)
    finally:
        wb.close()
    return Converted(
        markdown="\n".join(parts),
        converter="xlsx (openpyxl)",
        kind_hint="spreadsheet",
        meta={"sheet_count": sheet_count},
    ), "\n".join(ref)


def _pptx(path: Path, _attach: Path) -> tuple[Converted, str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts, ref = [f"# {path.stem}\n"], []
    for n, slide in enumerate(prs.slides, 1):
        title = (
            slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else ""
        )
        parts.append(f"## Slide {n}" + (f": {title}" if title else "") + "\n")
        ref.append(title)
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip() + "\n")
                ref.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                if not rows:
                    continue
                parts.append(md_table(rows[0], rows[1:]) + "\n")
                ref.extend(" ".join(r) for r in rows)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes = slide.notes_slide.notes_text_frame.text.strip()
            parts.append(f"**Speaker notes:** {notes}\n")
            ref.append(notes)
    return Converted(
        markdown="\n".join(parts),
        converter="pptx (python-pptx)",
        kind_hint="presentation",
        meta={"slide_count": len(prs.slides)},
    ), "\n".join(ref)


# ---- email -------------------------------------------------------------------------


def _email_markdown(subject, sender, to, cc, date, attachments, body_md) -> str:
    def fmt(people):
        return ", ".join(f"{n} <{a}>" if n and a else (n or a) for n, a in people)

    rows = [["From", fmt(sender)], ["To", fmt(to)]]
    if cc:
        rows.append(["Cc", fmt(cc)])
    if date:
        rows.append(["Date", date])
    if attachments:
        rows.append(["Attachments", ", ".join(attachments)])
    header = md_table(["Field", "Value"], rows)
    return f"# {subject or '(no subject)'}\n\n{header}\n\n## Message\n\n{body_md}\n"


def _save_attachment(attach: Path, filename: str, data: bytes) -> Path:
    """Save one attachment under a safe, unique file name inside ``attach``."""
    attach.mkdir(parents=True, exist_ok=True)
    # No path separators, and no leading dots (".", ".." or hidden files).
    safe = re.sub(r"[^\w.\- ]", "_", filename).strip().lstrip(".") or "attachment"
    out, n = attach / safe, 2
    while out.exists():  # two attachments with the same name must not overwrite each other
        out = attach / f"{Path(safe).stem}-{n}{Path(safe).suffix}"
        n += 1
    out.write_bytes(data)
    return out


def _eml(path: Path, attach: Path) -> tuple[Converted, str]:
    msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    subject = str(msg.get("subject", "") or "")
    sender = getaddresses([str(msg.get("from", ""))])
    to = getaddresses([str(v) for v in msg.get_all("to", [])])
    cc = getaddresses([str(v) for v in msg.get_all("cc", [])])
    date = ""
    if msg.get("date"):
        try:
            date = parsedate_to_datetime(str(msg["date"])).isoformat()
        except (TypeError, ValueError):
            date = str(msg["date"])
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body, body_ref = "", ""
    if body_part is not None:
        content = body_part.get_content()
        if body_part.get_content_subtype() == "html":
            body, body_ref = (
                html_to_markdown(content),
                BeautifulSoup(content, "html.parser").get_text(" "),
            )
        else:
            body = body_ref = content
    saved, names = [], []
    for part in msg.iter_attachments():
        name = part.get_filename() or "attachment"
        names.append(name)
        saved.append(_save_attachment(attach, name, part.get_payload(decode=True) or b""))
    conv = Converted(
        markdown=_email_markdown(subject, sender, to, cc, date, names, body),
        converter="eml (email stdlib)",
        kind_hint="email",
        attachments=saved,
        meta={
            "subject": subject,
            "from": sender,
            "to": to,
            "cc": cc,
            "date": date,
            "message_id": str(msg.get("message-id", "") or ""),
            "attachments": names,
        },
    )
    ref = " ".join([subject, *(n + " " + a for n, a in sender + to + cc), body_ref, *names])
    return conv, ref


def _msg(path: Path, attach: Path) -> tuple[Converted, str]:
    import extract_msg

    m = extract_msg.openMsg(str(path))
    try:
        sender = getaddresses([m.sender or ""])
        to = getaddresses([m.to or ""])
        cc = getaddresses([m.cc or ""])
        date = m.date.isoformat() if hasattr(m.date, "isoformat") else str(m.date or "")
        body = m.body or ""
        saved, names = [], []
        for a in m.attachments:
            name = a.longFilename or a.shortFilename or "attachment"
            if isinstance(a.data, bytes):
                names.append(name)
                saved.append(_save_attachment(attach, name, a.data))
        subject = m.subject or ""
    finally:
        m.close()
    conv = Converted(
        markdown=_email_markdown(subject, sender, to, cc, date, names, body),
        converter="msg (extract-msg)",
        kind_hint="email",
        attachments=saved,
        meta={
            "subject": subject,
            "from": sender,
            "to": to,
            "cc": cc,
            "date": date,
            "attachments": names,
        },
    )
    ref = " ".join([subject, *(n + " " + a for n, a in sender + to + cc), body, *names])
    return conv, ref


# ---- subtitles / meeting transcripts ----------------------------------------------------

VOICE_RE = re.compile(r"<v(?:\.[^ >]*)?\s+([^>]+)>(.*?)(?:</v>|$)", re.S)
TAG_RE = re.compile(r"<[^>]+>")


def _subtitles(path: Path, _attach: Path) -> tuple[Converted, str]:
    text = read_text_any(path)
    lines_out, ref, speakers = (
        [f"# {path.stem.replace('_', ' ').replace('-', ' ').title()}", ""],
        [],
        Counter(),
    )
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        rows = [r for r in block.strip().splitlines() if r.strip()]
        timing = next((r for r in rows if "-->" in r), None)
        if not timing:
            continue
        start = timing.split("-->")[0].strip().split(".")[0].split(",")[0]
        cue = " ".join(rows[rows.index(timing) + 1 :])
        m = VOICE_RE.search(cue)
        if m:
            speaker, said = m.group(1).strip(), TAG_RE.sub("", m.group(2)).strip()
        else:
            plain = TAG_RE.sub("", cue).strip()
            sm = re.match(r"^([A-Z][A-Za-z.'\- ]{1,40}):\s+(.*)$", plain)
            speaker, said = (sm.group(1), sm.group(2)) if sm else ("", plain)
        if not said:
            continue
        if speaker:
            speakers[speaker] += 1
            lines_out.append(f"**{speaker}** [{start}]: {said}")
            ref.append(f"{speaker} {said}")
        else:
            lines_out.append(f"[{start}] {said}")
            ref.append(said)
    conv = Converted(
        markdown="\n".join(lines_out) + "\n",
        converter=f"{path.suffix[1:]} (subtitle parser)",
        kind_hint="transcript",
        meta={"participants": [s for s, _ in speakers.most_common()]},
    )
    return conv, "\n".join(ref)


def attachment_key(parent_key: str, path: Path) -> str:
    # Single hyphen: gbrain collapses "--" in slugs.
    """Document key for an attachment: the parent's key plus the file name."""
    return f"{parent_key}-{slugify(path.stem, 40)}"
